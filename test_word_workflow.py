import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
import database
import workflow_store as store
from word_workflow import read_word,validate_fields,validate_values,parse_page_json,csv_bytes


FIELDS=[{'key':'species','label':'魚種','scope':'detail','type':'text'},
        {'key':'weight','label':'重量','scope':'detail','type':'decimal'}]


class WordWorkflowTests(unittest.TestCase):
    def test_docx_table_and_doc_rejection(self):
        doc=Document(); doc.add_paragraph('測試母版'); table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='重量'; table.cell(0,1).text='尾數'
        buffer=io.BytesIO(); doc.save(buffer)
        self.assertEqual(read_word(buffer.getvalue(),'x.docx')['tables'][0][0],['重量','尾數'])
        with self.assertRaises(ValueError): read_word(buffer.getvalue(),'x.doc')

    def test_unknown_keys_duplicates_numbers_and_zero(self):
        self.assertEqual(validate_values({'species':'甲','weight':0},FIELDS)['weight'],0)
        self.assertIsNone(validate_values({'species':'甲'},FIELDS)['weight'])
        for values in [{'wrong':1},{'weight':'NaN'},{'weight':True}]:
            with self.assertRaises(ValueError): validate_values(values,FIELDS)
        with self.assertRaises(ValueError): validate_fields(FIELDS+FIELDS)

    def test_no_truncation_repair_or_implicit_empty_success(self):
        for raw in ['{"matched":true,"rows":[', '{"matched":true,"rows":[]}', '{"matched":false,"rows":[]}']:
            with self.assertRaises(ValueError): parse_page_json(raw,FIELDS)
        self.assertEqual(parse_page_json('{"matched":true,"empty":true,"rows":[]}',FIELDS)['rows'],[])

    def test_csv_formula_safety_and_bom(self):
        data=csv_bytes([{'species':'=danger','weight':-1}],FIELDS)
        self.assertTrue(data.startswith(b'\xef\xbb\xbf'))
        self.assertIn("'=danger,-1",data.decode('utf-8-sig'))


class WorkflowPersistenceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='wf-test-',dir=Path.cwd())
        self.path=Path(self.temp.name)/'db.sqlite'
        def connect():
            conn=sqlite3.connect(self.path); conn.row_factory=sqlite3.Row
            conn.execute('PRAGMA foreign_keys=ON'); return conn
        self.connect=connect
        conn=connect(); conn.executescript(Path('migrations/002_word_workflow.sql').read_text(encoding='utf-8')); conn.close()
        self.patches=[patch.object(database,'get_db_connection',connect),patch.object(database,'IS_POSTGRES',False)]
        for p in self.patches: p.start()
        self.tid=store.publish('TEST','測試','t.docx',b'word',{'tables':[]},FIELDS)

    def tearDown(self):
        for p in reversed(self.patches): p.stop()
        self.temp.cleanup()

    def test_version_and_content_deduplication(self):
        self.assertEqual(self.tid,store.publish('TEST','測試','renamed.docx',b'word',{'tables':[]},FIELDS))
        store.publish('TEST','測試新版','t.docx',b'word2',{'tables':[]},FIELDS)
        self.assertEqual([x['version'] for x in store.templates()],[2,1])
        uid=store.register(self.tid,'a.pdf',b'pdf',1)
        self.assertEqual(uid,store.register(self.tid,'renamed.pdf',b'pdf',1))

    def test_all_pages_required_atomic_rollback_and_no_blank_rows(self):
        uid=store.register(self.tid,'a.pdf',b'pdf',2)
        store.save_page(uid,1,{'rows':[],'warnings':''},'manual','manual')
        with self.assertRaises(ValueError): store.confirm(uid,[{'source_page':1,'species':'甲'}],FIELDS)
        store.save_page(uid,2,{'rows':[],'warnings':''},'manual','manual')
        with self.assertRaises(ValueError):
            store.confirm(uid,[{'source_page':1,'species':'甲'},{'source_page':2,'weight':'bad'}],FIELDS)
        self.assertEqual(store.export_rows(self.tid),[])
        with self.assertRaises(ValueError): store.confirm(uid,[{'source_page':1}],FIELDS)

    def test_multiple_uploads_accumulate_and_survive_reopen(self):
        for index in range(2):
            uid=store.register(self.tid,str(index)+'.pdf',str(index).encode(),1)
            store.save_page(uid,1,{'rows':[{'species':'甲','weight':index}],'warnings':''},'manual','manual')
            self.assertTrue(store.confirm(uid,[{'source_page':1,'species':'甲','weight':index}],FIELDS))
            self.assertFalse(store.confirm(uid,[{'source_page':1,'species':'甲'}],FIELDS))
            with self.assertRaises(ValueError): store.save_page(uid,1,None,'manual','manual')
        self.assertEqual(len(store.export_rows(self.tid)),2)
        conn=self.connect()
        self.assertEqual(conn.execute('SELECT COUNT(*) FROM wf_records').fetchone()[0],2)
        conn.close()


if __name__=='__main__': unittest.main()
